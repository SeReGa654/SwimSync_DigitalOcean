"""Generator for Named Application (.docx)."""

from io import BytesIO
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from utils.ms_format import ms_to_csv_format


def _style_short(style: str) -> str:
    value = (style or "").strip().lower()
    mapping = {
        "freestyle": "в/с",
        "free": "в/с",
        "вільний": "в/с",
        "вільний стиль": "в/с",
        "breaststroke": "брас",
        "breast": "брас",
        "брас": "брас",
        "backstroke": "н/с",
        "back": "н/с",
        "спина": "н/с",
        "на спині": "н/с",
        "butterfly": "бат",
        "fly": "бат",
        "батерфляй": "бат",
        "бат": "бат",
        "medley": "к/п",
        "individual medley": "к/п",
        "комплексне плавання": "к/п",
        "к/п": "к/п",
    }
    return mapping.get(value, style or "—")


def _rank_ua(raw_rank: str) -> str:
    rank = (raw_rank or "").strip()
    if not rank:
        return "—"
    upper = rank.upper()
    mapping = {
        "MSMK": "МСМК",
        "MS": "МС",
        "KMS": "КМС",
        "KMSU": "КМСУ",
        "R1": "І",
        "R2": "ІІ",
        "R3": "ІІІ",
        "Y1": "І юн.",
        "Y2": "ІІ юн.",
        "Y3": "ІІІ юн.",
        "NONE": "—",
    }
    if upper in mapping:
        return mapping[upper]
    return rank


def _format_distance_line(event: Dict[str, Any]) -> str:
    distance = event.get("distanceM")
    style = _style_short(str(event.get("style") or ""))
    if not distance:
        return "—"
    return f"{distance} {style}"


def _format_time_line(event: Dict[str, Any]) -> str:
    value = event.get("entryTimeMs")
    if value is None:
        return "NT"
    return ms_to_csv_format(value)


def generate_named_application(payload: Dict[str, Any]) -> bytes:
    title = payload.get("title") or "Іменна заявка"
    organization = payload.get("organization") or "—"
    region = payload.get("region") or "—"
    generated_at = payload.get("generatedAt") or payload.get("generated_at") or "—"
    athletes: List[Dict[str, Any]] = payload.get("athletes") or []

    document = Document()
    heading = document.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(f"Спортивна школа/організація: {organization}")
    document.add_paragraph(f"Регіон: {region}")
    document.add_paragraph(f"Дата формування: {generated_at}")
    document.add_paragraph("")

    table = document.add_table(rows=1, cols=10)
    table.style = "Table Grid"
    headers = table.rows[0].cells
    headers[0].text = "№"
    headers[1].text = "Прізвище та ім'я спортсмена"
    headers[2].text = "Дата народження"
    headers[3].text = "Звання, розряд"
    headers[4].text = "Дистанція"
    headers[5].text = "Попередній час"
    headers[6].text = "Спортивна школа"
    headers[7].text = "Регіон, ОТГ"
    headers[8].text = "Прізвище, ініціали тренера"
    headers[9].text = "Віза лікаря про допуск"

    for row_index, athlete in enumerate(athletes, start=1):
        row = table.add_row().cells
        events = athlete.get("events") or []
        distance_lines = [_format_distance_line(event) for event in events] or ["—"]
        time_lines = [_format_time_line(event) for event in events] or ["NT"]

        row[0].text = str(row_index)
        row[1].text = athlete.get("fullName") or "—"
        row[2].text = str(athlete.get("birthYear") or "—")
        row[3].text = _rank_ua(str(athlete.get("rank") or ""))
        row[4].text = "\n".join(distance_lines)
        row[5].text = "\n".join(time_lines)
        row[6].text = athlete.get("club") or organization
        row[7].text = athlete.get("region") or region
        row[8].text = athlete.get("coach") or "—"
        row[9].text = ""

    output = BytesIO()
    document.save(output)
    return output.getvalue()
