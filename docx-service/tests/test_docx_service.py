from io import BytesIO
import time
import sys
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient

docx_root = str(Path(__file__).resolve().parents[1])
if docx_root not in sys.path:
    sys.path.insert(0, docx_root)
sys.modules.pop("parser", None)

from generator.start_protocol import generate_start_protocol
from generator.result_protocol import generate_result_protocol
from main import app


SAMPLE_COMPETITION = {
    "name": "Чемпіонат області",
    "categories_str": "Юнаки та дівчата",
    "location": "Київ",
    "venue": "Басейн ДЮСШ",
    "pool_length": 50,
    "date_from": "2026-05-10",
    "date_to": "2026-05-11",
}


SAMPLE_START_EVENTS = [
    {
        "distance_m": 50,
        "style": "Freestyle",
        "gender": "M",
        "heats": [
            {
                "number": 1,
                "entries": [
                    {
                        "lane": 4,
                        "full_name": "Іваненко Олексій",
                        "age_group": "2010-2011",
                        "birth_year": 2010,
                        "entry_time_ms": 28500,
                        "coach": "Петренко",
                    }
                ],
            }
        ],
    }
]


SAMPLE_RESULT_EVENTS = [
    {
        "distance_m": 50,
        "style": "Freestyle",
        "gender": "M",
        "age_groups": [
            {
                "name": "2010-2011",
                "results": [
                    {
                        "place": 1,
                        "place_display": "І",
                        "lane": 4,
                        "full_name": "Іваненко Олексій",
                        "birth_year": 2010,
                        "rank": "R1",
                        "club": "ДЮСШ-1",
                        "region": "Київ",
                        "finish_time_ms": 28000,
                        "achieved_rank": "R1",
                        "points_wa": 500,
                        "coach": "Петренко",
                        "status": "OK",
                        "dq_reason": None,
                    }
                ],
            }
        ],
    }
]

SAMPLE_NAMED_APPLICATION = {
    "title": "Іменна заявка",
    "organization": "ДЮСШ-1",
    "region": "Київ",
    "generatedAt": "2026-05-10",
    "athletes": [
        {
            "fullName": "Іваненко Олексій",
            "birthYear": 2010,
            "gender": "M",
            "rank": "R1",
            "coach": "Петренко",
            "club": "ДЮСШ-1",
            "region": "Київ",
            "events": [
                {"distanceM": 50, "style": "Freestyle", "entryTimeMs": 28500},
                {"distanceM": 100, "style": "Backstroke", "entryTimeMs": 64000},
            ],
        }
    ],
}


def _doc_text(docx_bytes: bytes) -> str:
    document = Document(BytesIO(docx_bytes))
    parts = [par.text for par in document.paragraphs if par.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def test_generate_start_protocol_golden_content():
    docx_bytes = generate_start_protocol(SAMPLE_COMPETITION, SAMPLE_START_EVENTS)
    text = _doc_text(docx_bytes)
    assert "Чемпіонат області" in text
    assert "Стартовий протокол" in text
    assert "День проведення: 2026-05-10" in text


def test_generate_result_protocol_golden_content():
    docx_bytes = generate_result_protocol(SAMPLE_COMPETITION, SAMPLE_RESULT_EVENTS)
    text = _doc_text(docx_bytes)
    assert "Чемпіонат області" in text
    assert "Протокол результатів" in text
    assert "50м Вільний стиль Чол. (2010-2011)" in text


def test_generate_named_application_endpoint():
    with TestClient(app) as client:
        response = client.post("/named-application", json=SAMPLE_NAMED_APPLICATION)
        assert response.status_code == 200
        assert (
            response.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        text = _doc_text(response.content)
        assert "Іменна заявка" in text
        assert "Іваненко Олексій" in text
        assert "Прізвище та ім'я спортсмена" in text
        assert "Звання, розряд" in text
        assert "І" in text
        assert "50 в/с" in text
        assert "100 н/с" in text

        document = Document(BytesIO(response.content))
        table = document.tables[0]
        assert len(table.rows) == 2  # header + 1 athlete row


def test_docx_queue_submit_status_download():
    with TestClient(app) as client:
        response = client.post(
            "/jobs/start-protocol",
            json={"competition": SAMPLE_COMPETITION, "events": SAMPLE_START_EVENTS},
        )
        assert response.status_code == 200
        job_id = response.json()["job_id"]

        status = "queued"
        for _ in range(40):
            status_resp = client.get(f"/jobs/{job_id}")
            assert status_resp.status_code == 200
            status = status_resp.json()["status"]
            if status == "completed":
                break
            if status == "failed":
                raise AssertionError("DOCX queue job failed")
            time.sleep(0.05)

        assert status == "completed"
        download_resp = client.get(f"/jobs/{job_id}/download")
        assert download_resp.status_code == 200
        assert (
            download_resp.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


def test_docx_queue_idempotency_key_reuses_job():
    with TestClient(app) as client:
        payload = {
            "competition": SAMPLE_COMPETITION,
            "events": SAMPLE_START_EVENTS,
            "idempotency_key": "same-request-key",
        }
        first = client.post("/jobs/start-protocol", json=payload)
        second = client.post("/jobs/start-protocol", json=payload)
        assert first.status_code == 200
        assert second.status_code == 200
        assert first.json()["job_id"] == second.json()["job_id"]
